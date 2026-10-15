"""
Экспорт расписания в Excel (.xlsx).
Использует openpyxl.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

# Цвета для предметов (те же, что в desktop_app.py)
SUBJECT_COLORS = [
    "6366F1", "EF4444", "10B981", "F59E0B",
    "8B5CF6", "06B6D4", "EC4899", "84CC16",
    "F97316", "14B8A6", "A855F7", "EAB308",
]

DAY_NAMES = ["Пн", "Вт", "Ср", "Чт", "Пт"]


def subject_color_hex(subject_name, subjects_list):
    """
    Возвращает HEX-цвет (без #) для предмета.
    subjects_list — список словарей вида [{"name": "Математика", ...}, ...]
    """
    if not subject_name:
        return SUBJECT_COLORS[0]
    for i, s in enumerate(subjects_list):
        if s.get("name") == subject_name:
            return SUBJECT_COLORS[i % len(SUBJECT_COLORS)]
    return "64748B"   # серый, если не нашли
def write_schedule_sheet(ws, class_name, schedule, subjects_list):
    """
    Пишет расписание одного класса на лист ws.
    schedule — список объектов Lesson
    subjects_list — список словарей предметов
    """
    # Заголовок
    ws.merge_cells("A1:F1")
    ws["A1"] = f"Расписание: {class_name}"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1E293B")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    # Шапка таблицы
    headers = ["Урок"] + DAY_NAMES
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style="thin", color="94A3B8"),
            right=Side(style="thin", color="94A3B8"),
            top=Side(style="thin", color="94A3B8"),
            bottom=Side(style="thin", color="94A3B8"),
        )

    # Сетка расписания
    for slot in range(7):
        row = slot + 3
        # Номер урока
        cell = ws.cell(row=row, column=1, value=slot + 1)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style="thin", color="94A3B8"),
            right=Side(style="thin", color="94A3B8"),
            top=Side(style="thin", color="94A3B8"),
            bottom=Side(style="thin", color="94A3B8"),
        )

        # Дни
        for day in range(5):
            col = day + 2
            # Ищем урок для этого класса в этот день/слот
            lesson = None
            for l in schedule:
                if l.class_name == class_name and l.day == day and l.slot == slot:
                    lesson = l
                    break

            cell = ws.cell(row=row, column=col)
            if lesson:
                cell.value = f"{lesson.subject}\nкаб. {lesson.room_id}"
                color = subject_color_hex(lesson.subject, subjects_list)
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(bold=True, color="FFFFFF", size=10)
            else:
                cell.value = "—"
                cell.font = Font(color="94A3B8")

            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = Border(
                left=Side(style="thin", color="94A3B8"),
                right=Side(style="thin", color="94A3B8"),
                top=Side(style="thin", color="94A3B8"),
                bottom=Side(style="thin", color="94A3B8"),
            )
            ws.row_dimensions[row].height = 35

    # Ширина столбцов
    ws.column_dimensions["A"].width = 8
    for col in range(2, 7):
        ws.column_dimensions[get_column_letter(col)].width = 22

def write_teachers_sheet(ws, schedule, teachers_list, subjects_list):
    """
    Пишет расписание ВСЕХ учителей на один лист ws.
    """
    if not teachers_list:
        ws["A1"] = "Нет учителей для экспорта"
        return

    # Собираем id учителей, у которых есть уроки
    teacher_ids_with_lessons = set()
    for lesson in schedule:
        teacher_ids_with_lessons.add(lesson.teacher_id)

    current_row = 1

    for teacher in teachers_list:
        t_id = teacher.get("id")
        t_name = teacher.get("name", "?")

        if t_id not in teacher_ids_with_lessons:
            continue

        # Заголовок
        ws.merge_cells(
            start_row=current_row, start_column=1,
            end_row=current_row, end_column=6
        )
        title_cell = ws.cell(row=current_row, column=1)
        title_cell.value = f"Учитель: {t_name}"
        title_cell.font = Font(bold=True, size=14, color="FFFFFF")
        title_cell.fill = PatternFill("solid", fgColor="1E3A5F")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[current_row].height = 25
        current_row += 1

        # Шапка
        headers = ["Урок"] + DAY_NAMES
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=current_row, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334155")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                left=Side(style="thin", color="94A3B8"),
                right=Side(style="thin", color="94A3B8"),
                top=Side(style="thin", color="94A3B8"),
                bottom=Side(style="thin", color="94A3B8"),
            )
        current_row += 1

        # Сетка 7×5
        for slot in range(7):
            cell = ws.cell(row=current_row, column=1, value=slot + 1)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                left=Side(style="thin", color="94A3B8"),
                right=Side(style="thin", color="94A3B8"),
                top=Side(style="thin", color="94A3B8"),
                bottom=Side(style="thin", color="94A3B8"),
            )

            for day in range(5):
                col = day + 2
                lesson = None
                for l in schedule:
                    if l.teacher_id == t_id and l.day == day and l.slot == slot:
                        lesson = l
                        break

                cell = ws.cell(row=current_row, column=col)
                if lesson:
                    cell.value = f"{lesson.class_name}: {lesson.subject}"
                    color = subject_color_hex(lesson.subject, subjects_list)
                    cell.fill = PatternFill("solid", fgColor=color)
                    cell.font = Font(bold=True, color="FFFFFF", size=10)
                else:
                    cell.value = "—"
                    cell.font = Font(color="94A3B8")

                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
                cell.border = Border(
                    left=Side(style="thin", color="94A3B8"),
                    right=Side(style="thin", color="94A3B8"),
                    top=Side(style="thin", color="94A3B8"),
                    bottom=Side(style="thin", color="94A3B8"),
                )
            ws.row_dimensions[current_row].height = 35
            current_row += 1

        current_row += 2

    # Ширина столбцов
    ws.column_dimensions["A"].width = 8
    for col in range(2, 7):
        ws.column_dimensions[get_column_letter(col)].width = 22
def write_teachers_data_sheet(ws, teachers_list):
    """Список учителей с их предметами."""
    headers = ["ID", "ФИО", "Предметы"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style="thin", color="94A3B8"),
            right=Side(style="thin", color="94A3B8"),
            top=Side(style="thin", color="94A3B8"),
            bottom=Side(style="thin", color="94A3B8"),
        )

    for row_idx, t in enumerate(teachers_list, start=2):
        ws.cell(row=row_idx, column=1, value=t.get("id", ""))
        ws.cell(row=row_idx, column=2, value=t.get("name", ""))
        ws.cell(row=row_idx, column=3, value=", ".join(t.get("subjects", [])))
        for col in range(1, 4):
            cell = ws.cell(row=row_idx, column=col)
            cell.border = Border(
                left=Side(style="thin", color="94A3B8"),
                right=Side(style="thin", color="94A3B8"),
                top=Side(style="thin", color="94A3B8"),
                bottom=Side(style="thin", color="94A3B8"),
            )

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 40


def write_classes_data_sheet(ws, classes_list):
    """Список классов с числом учеников."""
    headers = ["ID", "Название", "Учеников"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style="thin", color="94A3B8"),
            right=Side(style="thin", color="94A3B8"),
            top=Side(style="thin", color="94A3B8"),
            bottom=Side(style="thin", color="94A3B8"),
        )

    for row_idx, c in enumerate(classes_list, start=2):
        ws.cell(row=row_idx, column=1, value=c.get("id", ""))
        ws.cell(row=row_idx, column=2, value=c.get("name", ""))
        ws.cell(row=row_idx, column=3, value=c.get("students", ""))
        for col in range(1, 4):
            cell = ws.cell(row=row_idx, column=col)
            cell.border = Border(
                left=Side(style="thin", color="94A3B8"),
                right=Side(style="thin", color="94A3B8"),
                top=Side(style="thin", color="94A3B8"),
                bottom=Side(style="thin", color="94A3B8"),
            )

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 12


def write_subjects_data_sheet(ws, subjects_list, classes_list):
    """Предметы и часы по классам."""
    headers = ["Предмет"] + [c.get("name", "?") for c in classes_list]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style="thin", color="94A3B8"),
            right=Side(style="thin", color="94A3B8"),
            top=Side(style="thin", color="94A3B8"),
            bottom=Side(style="thin", color="94A3B8"),
        )

    for row_idx, s in enumerate(subjects_list, start=2):
        ws.cell(row=row_idx, column=1, value=s.get("name", ""))
        hours = s.get("hours", {})
        for col_idx, c in enumerate(classes_list, start=2):
            cname = c.get("name", "")
            ws.cell(row=row_idx, column=col_idx, value=hours.get(cname, 0))
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                left=Side(style="thin", color="94A3B8"),
                right=Side(style="thin", color="94A3B8"),
                top=Side(style="thin", color="94A3B8"),
                bottom=Side(style="thin", color="94A3B8"),
            )

    ws.column_dimensions["A"].width = 25
    for col in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 10


def write_rooms_data_sheet(ws, rooms_list):
    """Список кабинетов."""
    headers = ["ID", "Название", "Мест", "Спец. для предмета"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style="thin", color="94A3B8"),
            right=Side(style="thin", color="94A3B8"),
            top=Side(style="thin", color="94A3B8"),
            bottom=Side(style="thin", color="94A3B8"),
        )

    for row_idx, r in enumerate(rooms_list, start=2):
        ws.cell(row=row_idx, column=1, value=r.get("id", ""))
        ws.cell(row=row_idx, column=2, value=r.get("name", ""))
        ws.cell(row=row_idx, column=3, value=r.get("capacity", ""))
        ws.cell(row=row_idx, column=4, value=r.get("subject_specific", ""))
        for col in range(1, 5):
            cell = ws.cell(row=row_idx, column=col)
            cell.border = Border(
                left=Side(style="thin", color="94A3B8"),
                right=Side(style="thin", color="94A3B8"),
                top=Side(style="thin", color="94A3B8"),
                bottom=Side(style="thin", color="94A3B8"),
            )

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 25
def export_to_excel(schedule, data, filepath):
    """
    Главная функция экспорта.
    Создаёт Excel-файл с вкладками:
      - Расписание <класс> (для каждого класса)
      - По учителям
      - Учителя
      - Классы
      - Предметы
      - Кабинеты
    Возвращает (успех: bool, сообщение: str).
    """
    try:
        wb = Workbook()

        # Удаляем автоматически созданный лист "Sheet"
        default_sheet = wb.active
        wb.remove(default_sheet)

        classes_list = data.get("classes", [])
        teachers_list = data.get("teachers", [])
        subjects_list = data.get("subjects", [])
        rooms_list = data.get("rooms", [])

        # --- Лист на каждый класс ---
        for cls in classes_list:
            cname = cls.get("name", "?")
            # Excel-листы не любят некоторые символы в имени
            sheet_name = f"Расписание {cname}"[:31]   # лимит 31 символ
            ws = wb.create_sheet(title=sheet_name)
            write_schedule_sheet(ws, cname, schedule, subjects_list)

        # --- Лист «По учителям» ---
        ws_teachers = wb.create_sheet(title="По учителям")
        write_teachers_sheet(ws_teachers, schedule, teachers_list, subjects_list)

        # --- Лист «Учителя» ---
        ws_t_data = wb.create_sheet(title="Учителя")
        write_teachers_data_sheet(ws_t_data, teachers_list)

        # --- Лист «Классы» ---
        ws_c_data = wb.create_sheet(title="Классы")
        write_classes_data_sheet(ws_c_data, classes_list)

        # --- Лист «Предметы» ---
        ws_s_data = wb.create_sheet(title="Предметы")
        write_subjects_data_sheet(ws_s_data, subjects_list, classes_list)

        # --- Лист «Кабинеты» ---
        ws_r_data = wb.create_sheet(title="Кабинеты")
        write_rooms_data_sheet(ws_r_data, rooms_list)

        # Сохраняем
        wb.save(filepath)
        return True, f"Экспортировано в {filepath}"

    except PermissionError:
        return False, ("Не удалось сохранить файл. Возможно, он открыт в Excel. "
                       "Закройте его и попробуйте снова.")
    except Exception as e:
        return False, f"Ошибка экспорта: {e}"